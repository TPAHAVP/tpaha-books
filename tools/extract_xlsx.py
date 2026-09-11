"""Dump everything useful from the two TPAHA workbooks to JSON for reimplementation + test oracles."""
import json, sys, re, os
from openpyxl import load_workbook

def rgb(fill):
    try:
        if fill is None or fill.fill_type is None: return None
        c = fill.fgColor
        if c is None: return None
        if c.type == 'rgb' and c.rgb and c.rgb not in ('00000000',): return c.rgb
        if c.type == 'theme': return f"theme{c.theme}{'+'+str(round(c.tint,2)) if c.tint else ''}"
        if c.type == 'indexed': return f"idx{c.indexed}"
    except Exception: return None
    return None

def dump(path, out):
    wbf = load_workbook(path, data_only=False)
    wbv = load_workbook(path, data_only=True)
    res = {"file": os.path.basename(path), "sheets": {}, "definedNames": {}, "tables": {}}
    for name, dn in wbf.defined_names.items():
        res["definedNames"][name] = dn.attr_text
    for ws in wbf.worksheets:
        wv = wbv[ws.title]
        sh = {"state": ws.sheet_state, "dims": ws.dimensions, "max_row": ws.max_row, "max_col": ws.max_column,
              "cells": {}, "merged": [str(r) for r in ws.merged_cells.ranges], "dv": [], "cf": [],
              "colWidths": {}, "rowHeights": {}, "tables": {}, "protection": bool(ws.protection.sheet)}
        for row in ws.iter_rows():
            for c in row:
                v = wv[c.coordinate].value
                f = c.value if isinstance(c.value, str) and c.value.startswith('=') else None
                if c.value is None and v is None and rgb(c.fill) is None: continue
                ent = {}
                if f: ent["f"] = f
                if v is not None: ent["v"] = v if not hasattr(v, 'isoformat') else v.isoformat()
                if c.number_format and c.number_format != 'General': ent["fmt"] = c.number_format
                fl = rgb(c.fill)
                if fl: ent["fill"] = fl
                if c.font and c.font.bold: ent["b"] = True
                if c.alignment and c.alignment.wrap_text: ent["wrap"] = True
                if c.protection and not c.protection.locked: ent["unlocked"] = True
                if ent: sh["cells"][c.coordinate] = ent
        for dv in ws.data_validations.dataValidation:
            sh["dv"].append({"sqref": str(dv.sqref), "type": dv.type, "op": dv.operator, "f1": dv.formula1, "f2": dv.formula2,
                             "prompt": dv.prompt, "error": dv.error, "allowBlank": dv.allowBlank})
        for cf in ws.conditional_formatting:
            for r in cf.rules:
                sh["cf"].append({"sqref": str(cf.sqref), "type": r.type, "op": r.operator, "text": r.text, "formula": list(r.formula) if r.formula else None,
                                 "fill": rgb(r.dxf.fill) if r.dxf and r.dxf.fill else None,
                                 "font": (r.dxf.font.color.rgb if r.dxf and r.dxf.font and r.dxf.font.color else None)})
        for k, d in ws.column_dimensions.items():
            if d.width: sh["colWidths"][k] = d.width
        for k, d in ws.row_dimensions.items():
            if d.height: sh["rowHeights"][str(k)] = d.height
        for t in ws.tables.values():
            tname = t.name
            sh["tables"][tname] = {"ref": t.ref, "columns": [c.name for c in t.tableColumns]}
            res["tables"][tname] = {"sheet": ws.title, "ref": t.ref, "columns": [c.name for c in t.tableColumns]}
        res["sheets"][ws.title] = sh
    with open(out, "w", encoding="utf-8") as fh:
        json.dump(res, fh, indent=1, ensure_ascii=False, default=str)
    return res

P = sys.argv[1]; S = sys.argv[2]
a = dump(os.path.join(P, "TPAHA_2026 (1).xlsx"), os.path.join(S, "extract", "a.json"))
b = dump(os.path.join(P, "TPAHA_Winter_Storage (1).xlsx"), os.path.join(S, "extract", "b.json"))

# Office Script sources embedded in Winter Storage 'Button Setup'
bs = b["sheets"]["Button Setup"]["cells"]
names = {"A13": "CompleteStorageSwap", "A19": "AddReserveLine", "A25": "RemoveReserveLine", "A36": "ExportPrintablePDF"}
for addr, nm in names.items():
    txt = bs.get(addr, {}).get("v", "")
    if isinstance(txt, str) and "function main" in txt:
        with open(os.path.join(S, "extract", "scripts", nm + ".ts"), "w", encoding="utf-8") as fh: fh.write(txt)
        print(f"script {nm}: {len(txt)} chars from {addr}")
    else:
        print(f"script {nm}: NOT FOUND at {addr} (first 60 chars: {str(txt)[:60]!r})")

# Summaries
print("\n=== A sheets ===")
for n, sh in a["sheets"].items():
    nf = sum(1 for c in sh["cells"].values() if "f" in c)
    print(f"  {n:14s} state={sh['state']:8s} dims={sh['dims']:10s} cells={len(sh['cells'])} formulas={nf} merged={len(sh['merged'])} dv={len(sh['dv'])} cf={len(sh['cf'])} prot={sh['protection']}")
print("definedNames:", a["definedNames"])
print("\n=== B sheets ===")
for n, sh in b["sheets"].items():
    nf = sum(1 for c in sh["cells"].values() if "f" in c)
    print(f"  {n:20s} state={sh['state']:8s} dims={sh['dims']:10s} cells={len(sh['cells'])} formulas={nf} merged={len(sh['merged'])} dv={len(sh['dv'])} cf={len(sh['cf'])} prot={sh['protection']}")
print("definedNames:", b["definedNames"])
fills = {}
for addr, c in b["sheets"]["Photo_Style_Layout"]["cells"].items():
    if "fill" in c: fills[c["fill"]] = fills.get(c["fill"], 0) + 1
print("Photo_Style_Layout fill colours:", fills)
