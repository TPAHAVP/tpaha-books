"""Build the app seed files and the test fixtures from the two TPAHA workbooks.

Outputs
  data/ledger-2026.json            real seed for the Treasurer Ledger (git-ignored)
  data/winter-storage.json         real seed for Winter Storage (git-ignored)
  data/fixtures/ledger-fixture.json      anonymised ledger (descriptions replaced) - not committed
  data/fixtures/ledger-oracle.json       month and annual figures cached by Excel
  data/fixtures/storage-fixture.json     anonymised storage (names, phones, plates replaced)
  data/fixtures/storage-oracle.json      per-spot figures cached by Excel

Run from the project root:  python tools/build_seed.py
"""
import json, os, re, sys, hashlib
from datetime import datetime, date, timezone
from openpyxl import load_workbook


def utc_now():
    return datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%S.000Z")

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
LEDGER_XLSX = os.path.join(ROOT, "TPAHA_2026 (1).xlsx")
STORAGE_XLSX = os.path.join(ROOT, "TPAHA_Winter_Storage (1).xlsx")
DATA = os.path.join(ROOT, "data")
FIX = os.path.join(ROOT, "data", "fixtures")   # derived from real workbooks: git-ignored
os.makedirs(DATA, exist_ok=True); os.makedirs(FIX, exist_ok=True)

MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August",
          "September", "October", "November", "December"]
CATEGORIES = [
    {"id": 1, "name": "Membership", "type": "Deposit"},
    {"id": 2, "name": "Grub Box", "type": "Deposit"},
    {"id": 3, "name": "Arena Storage", "type": "Deposit"},
    {"id": 4, "name": "Misc Income", "type": "Deposit"},
    {"id": 5, "name": "MB Hydro", "type": "Withdrawal"},
    {"id": 6, "name": "Taxes/Ins", "type": "Withdrawal"},
    {"id": 8, "name": "Misc Exp", "type": "Withdrawal"},
    {"id": 9, "name": "Bank Chgs", "type": "Withdrawal"},
]
CAT_COLS = {"Membership": "F", "Grub Box": "G", "Arena Storage": "H", "Misc Income": "I",
            "MB Hydro": "J", "Taxes/Ins": "K", "Misc Exp": "L", "Bank Chgs": "M"}


def uid_for(*parts):
    """Deterministic short id so re-running the seed keeps the same uids."""
    return hashlib.sha1("|".join(str(p) for p in parts).encode("utf-8")).hexdigest()[:12]


def cents(v):
    return int(round(float(v) * 100))


def iso_date(v):
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    return str(v)[:10]


def clean(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


# ----------------------------------------------------------------------------- ledger
def build_ledger():
    wb = load_workbook(LEDGER_XLSX, data_only=True)
    entry = wb["ENTRY"]
    log = wb["LOG"]
    year = int(wb["ConfigHidden"]["A1"].value)
    prior = float(entry["B23"].value)
    rows = []
    for r in range(2, log.max_row + 1):
        tid, ts, d, typ, cat, amt, desc, chq, notes = [log.cell(row=r, column=c).value for c in range(1, 10)]
        if tid is None and d is None:
            continue
        rows.append(dict(legacyId=tid, ts=ts, date=iso_date(d), type=typ, category=cat,
                         amount=round(float(amt), 2), description=str(desc or "").strip(),
                         chequeNum="" if chq is None else str(clean(chq)), notes=str(notes or "")))
    rows.sort(key=lambda x: str(x["ts"]))  # entry order (the LOG is already in timestamp order)
    txns = []
    for i, x in enumerate(rows, start=1):
        u = uid_for("ledger", year, x["legacyId"], x["ts"])
        txns.append({
            "uid": u, "number": i, "legacyId": x["legacyId"], "date": x["date"], "type": x["type"],
            "category": x["category"], "amount": x["amount"], "description": x["description"],
            "chequeNum": x["chequeNum"], "notes": x["notes"], "createdAt": str(x["ts"]),
            "createdBy": "", "updatedAt": str(x["ts"]), "updatedBy": "", "deleted": False,
        })
    doc = {"schema": 1, "year": year, "priorYearBalance": prior, "categories": CATEGORIES,
           "nextNumber": len(txns) + 1, "transactions": txns,
           "updatedAt": utc_now(), "updatedBy": "seed"}

    # Oracle: what Excel cached for each month sheet and the annual sheet.
    oracle = {"year": year, "priorYearBalance": prior, "months": {}, "annual": {}}
    for m, name in enumerate(MONTHS, start=1):
        ws = wb[name]
        by_cat = {cat: round(float(ws[f"{col}34"].value or 0), 2) for cat, col in CAT_COLS.items()}
        oracle["months"][m] = {
            "amountTotal": round(float(ws["C34"].value or 0), 2), "byCategory": by_cat,
            "net": round(float(ws["O34"].value or 0), 2), "opening": round(float(ws["B37"].value), 2),
            "deposits": round(float(ws["B38"].value or 0), 2), "expenses": round(float(ws["B39"].value or 0), 2),
            "closing": round(float(ws["B40"].value), 2),
            "rowCount": sum(1 for r in range(4, 34) if ws[f"A{r}"].value not in (None, "")),
        }
    an = wb["Annual 2026"]
    cat_rows = {"Membership": 4, "Grub Box": 5, "Arena Storage": 6, "Misc Income": 7,
                "MB Hydro": 10, "Taxes/Ins": 11, "Misc Exp": 12, "Bank Chgs": 13}
    cols = "BCDEFGHIJKLM"
    oracle["annual"]["byCategoryByMonth"] = {cat: [round(float(an[f"{cols[i]}{r}"].value or 0), 2) for i in range(12)]
                                             for cat, r in cat_rows.items()}
    oracle["annual"]["yearTotalByCategory"] = {cat: round(float(an[f"N{r}"].value or 0), 2) for cat, r in cat_rows.items()}
    # Rows 8/14/16 (Total Income / Total Expenses / Net per month) are only partly cached in the
    # workbook: Excel Online never finished calculating rows 14 and 16 for Feb-Dec (the sheet has a
    # circular reference in N14). Derive them from the fully cached category rows and keep the raw
    # cached cells for the README report.
    bcm = oracle["annual"]["byCategoryByMonth"]
    inc_cats = ["Membership", "Grub Box", "Arena Storage", "Misc Income"]
    exp_cats = ["MB Hydro", "Taxes/Ins", "Misc Exp", "Bank Chgs"]
    oracle["annual"]["totalIncomeByMonth"] = [round(sum(bcm[c][i] for c in inc_cats), 2) for i in range(12)]
    oracle["annual"]["totalExpensesByMonth"] = [round(sum(bcm[c][i] for c in exp_cats), 2) for i in range(12)]
    oracle["annual"]["netByMonth"] = [round(oracle["annual"]["totalIncomeByMonth"][i] - oracle["annual"]["totalExpensesByMonth"][i], 2) for i in range(12)]
    oracle["annual"]["workbookCached"] = {
        "row8_totalIncome": [an[f"{cols[i]}8"].value for i in range(12)],
        "row14_totalExpenses": [an[f"{cols[i]}14"].value for i in range(12)],
        "row16_net": [an[f"{cols[i]}16"].value for i in range(12)],
        "N14_totalExpensesYear": an["N14"].value, "N16_netYear": an["N16"].value,
    }
    oracle["annual"]["yearTotalIncome"] = round(float(an["N8"].value or 0), 2)
    # N14 and N16 are wrong in the workbook (self-referencing sum; N8-N15). Record the correct values.
    exp = sum(oracle["annual"]["yearTotalByCategory"][c] for c in ["MB Hydro", "Taxes/Ins", "Misc Exp", "Bank Chgs"])
    oracle["annual"]["yearTotalExpensesCorrect"] = round(exp, 2)
    oracle["annual"]["yearNetCorrect"] = round(oracle["annual"]["yearTotalIncome"] - exp, 2)
    oracle["annual"]["workbookYearNet"] = an["N16"].value
    oracle["annual"]["opening"] = round(float(an["B18"].value), 2)
    oracle["annual"]["closing"] = round(float(an["B19"].value), 2)

    fixture = json.loads(json.dumps(doc))
    for i, t in enumerate(fixture["transactions"], start=1):
        t["description"] = f"Transaction {i}"
        t["notes"] = ""
    fixture["updatedBy"] = "fixture"
    return doc, fixture, oracle


# ---------------------------------------------------------------------------- storage
# Cell mapping copied from the workbook's Complete Storage Swap script: 12 entry cells per spot.
SPOT_CELLS = [
    ("R1C1", "B5", ["B7", "B10", "B13", "B16", "B18", "C18", "D18", "B20", "D20", "B22", "B24", "AF5"]),
    ("R1C2", "F5", ["F7", "F10", "F13", "F16", "F18", "G18", "H18", "F20", "H20", "F22", "F24", "AF6"]),
    ("R1C3", "J5", ["J7", "J10", "J13", "J16", "J18", "K18", "L18", "J20", "L20", "J22", "J24", "AF7"]),
    ("R1C4", "N5", ["N7", "N10", "N13", "N16", "N18", "O18", "P18", "N20", "P20", "N22", "N24", "AF8"]),
    ("R1C5", "R5", ["R7", "R10", "R13", "R16", "R18", "S18", "T18", "R20", "T20", "R22", "R24", "AF9"]),
    ("R2C1", "B28", ["B30", "B33", "B36", "B39", "B41", "C41", "D41", "B43", "D43", "B45", "B47", "AF10"]),
    ("R2C2", "F28", ["F30", "F33", "F36", "F39", "F41", "G41", "H41", "F43", "H43", "F45", "F47", "AF11"]),
    ("R2C3", "J28", ["J30", "J33", "J36", "J39", "J41", "K41", "L41", "J43", "L43", "J45", "J47", "AF12"]),
    ("R2C4", "N28", ["N30", "N33", "N36", "N39", "N41", "O41", "P41", "N43", "P43", "N45", "N47", "AF13"]),
    ("R2C5", "R28", ["R30", "R33", "R36", "R39", "R41", "S41", "T41", "R43", "T43", "R45", "R47", "AF14"]),
    ("R3C1", "B51", ["B53", "B56", "B59", "B62", "B64", "C64", "D64", "B66", "D66", "B68", "B70", "AF15"]),
    ("R3C2", "F51", ["F53", "F56", "F59", "F62", "F64", "G64", "H64", "F66", "H66", "F68", "F70", "AF16"]),
    ("R3C3", "J51", ["J53", "J56", "J59", "J62", "J64", "K64", "L64", "J66", "L66", "J68", "J70", "AF17"]),
    ("R3C4", "N51", ["N53", "N56", "N59", "N62", "N64", "O64", "P64", "N66", "P66", "N68", "N70", "AF18"]),
    ("R3C5", "R51", ["R53", "R56", "R59", "R62", "R64", "S64", "T64", "R66", "T66", "R68", "R70", "AF19"]),
    ("R4C1", "B74", ["B76", "B79", "B82", "B85", "B87", "C87", "D87", "B89", "D89", "B91", "B93", "AF20"]),
    ("R4C2", "F74", ["F76", "F79", "F82", "F85", "F87", "G87", "H87", "F89", "H89", "F91", "F93", "AF21"]),
    ("R4C3", "J74", ["J76", "J79", "J82", "J85", "J87", "K87", "L87", "J89", "L89", "J91", "J93", "AF22"]),
    ("R4C4", "N74", ["N76", "N79", "N82", "N85", "N87", "O87", "P87", "N89", "P89", "N91", "N93", "AF23"]),
    ("R4C5", "R74", ["R76", "R79", "R82", "R85", "R87", "S87", "T87", "R89", "T89", "R91", "R93", "AF24"]),
    ("R5C1", "B97", ["B99", "B102", "B105", "B108", "B110", "C110", "D110", "B112", "D112", "B114", "B116", "AF25"]),
    ("R5C2", "F97", ["F99", "F102", "F105", "F108", "F110", "G110", "H110", "F112", "H112", "F114", "F116", "AF26"]),
    ("R5C3", "J97", ["J99", "J102", "J105", "J108", "J110", "K110", "L110", "J112", "L112", "J114", "J116", "AF27"]),
    ("R5C4", "N97", ["N99", "N102", "N105", "N108", "N110", "O110", "P110", "N112", "P112", "N114", "N116", "AF28"]),
    ("R5C5", "R97", ["R99", "R102", "R105", "R108", "R110", "S110", "T110", "R112", "T112", "R114", "R116", "AF29"]),
    ("END-1", "B124", ["B126", "B129", "B132", "B135", "B137", "E137", "H137", "B139", "H139", "B141", "B143", "AF30"]),
    ("END-2", "L124", ["L126", "L129", "L132", "L135", "L137", "O137", "R137", "L139", "R139", "L141", "L143", "AF31"]),
]
FIELDS = ["name", "equipment", "phone", "plate", "length", "width", "qty", "firstYear", "type", "feeOverride", "notes", "sourceRow"]
NUMERIC = {"length", "width", "qty", "firstYear", "feeOverride", "sourceRow"}
# Status-line cell under each main-row spot block, and the AG/AH/AI/AJ hidden calc row.
STATUS_CELL = {}
for idx, (spot, header, cells) in enumerate(SPOT_CELLS):
    col = re.match(r"[A-Z]+", header).group(0); row = int(re.search(r"\d+", header).group(0))
    STATUS_CELL[spot] = f"{col}{row + 21}" if not spot.startswith("END") else f"{col}{row + 21}"
CALC_ROW = {spot: 5 + i for i, (spot, _, _) in enumerate(SPOT_CELLS)}  # AF5..AF31
FOOTER_CELLS = {1: "B121", 2: "F121", 3: "J121", 4: "N121", 5: "R121"}


def num_or_none(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return int(v) if float(v).is_integer() else float(v)
    try:
        f = float(str(v).strip())
        return int(f) if f.is_integer() else f
    except ValueError:
        return str(v)


def build_storage():
    wb = load_workbook(STORAGE_XLSX, data_only=True)
    ws = wb["Photo_Style_Layout"]
    sl = wb["Storage List"]
    rl = wb["Reserve List"]
    now = utc_now()
    spots = []
    for spot, header, cells in SPOT_CELLS:
        assert ws[header].value == spot, (spot, header, ws[header].value)
        rec = {"spot": spot}
        for f, addr in zip(FIELDS, cells):
            v = ws[addr].value
            if f in NUMERIC:
                rec[f] = num_or_none(v)
            else:
                rec[f] = "" if v is None else str(v).strip()
        rec["updatedAt"] = now; rec["updatedBy"] = "seed"
        spots.append(rec)
    doc = {
        "schema": 1, "storageYear": int(ws["W5"].value),
        "rates": {"Boat": num_or_none(ws["AA10"].value), "Pontoon / Camper": num_or_none(ws["AA11"].value)},
        "arena": {"length": num_or_none(ws["AA38"].value), "width": num_or_none(ws["AA39"].value),
                  "firstRowBlockDepth": num_or_none(ws["AA40"].value), "gapBetweenRows": num_or_none(ws["AA41"].value),
                  "mainColumnWidth": num_or_none(ws["AA42"].value), "leftMargin": num_or_none(ws["AA43"].value),
                  "rightMargin": num_or_none(ws["AA44"].value), "gapBetweenColumns": num_or_none(ws["AA45"].value)},
        "endSpaces": {"END-1": {"length": num_or_none(ws["Y53"].value), "depth": num_or_none(ws["AA53"].value)},
                      "END-2": {"length": num_or_none(ws["Y54"].value), "depth": num_or_none(ws["AA54"].value)}},
        "spots": spots, "reserve": [], "updatedAt": now, "updatedBy": "seed",
    }
    for r in range(5, rl.max_row + 1):
        vals = [rl.cell(row=r, column=c).value for c in range(1, 9)]
        if all(v in (None, "") for v in vals):
            continue
        name, years, phone, equip, qty, length, notes, plate = vals
        doc["reserve"].append({
            "uid": uid_for("reserve", r, name), "name": str(name or "").strip(), "yearsNote": "" if years is None else str(clean(years)),
            "phone": "" if phone is None else str(phone).strip(), "equipment": str(equip or "").strip(),
            "qty": num_or_none(qty), "length": num_or_none(length), "notes": str(notes or "").strip(),
            "plate": "" if plate is None else str(plate).strip(), "deleted": False, "updatedAt": now, "updatedBy": "seed",
        })

    # Oracle from the Storage List sheet and the layout's status lines / footers / header.
    oracle = {"spots": {}, "footers": {}, "header": {}}
    for r in range(5, 32):
        spot = sl[f"A{r}"].value
        oracle["spots"][spot] = {
            "yearsStored": clean(sl[f"G{r}"].value), "unitFee": clean(sl[f"M{r}"].value), "totalFee": clean(sl[f"N{r}"].value),
            "customerTotal": clean(sl[f"O{r}"].value), "start": clean(sl[f"R{r}"].value), "end": clean(sl[f"S{r}"].value),
            "sizeCheck": clean(sl[f"T{r}"].value), "statusLine": clean(ws[STATUS_CELL[spot]].value),
        }
    for c, addr in FOOTER_CELLS.items():
        oracle["footers"][c] = ws[addr].value
    oracle["header"] = {"assigned": ws["W28"].value, "revenue": ws["W31"].value, "widthRemaining": ws["W47"].value}

    # Anonymised fixture: consistent aliases so duplicate customers stay duplicates.
    fixture = json.loads(json.dumps(doc))
    alias = {}
    def alias_for(name):
        if not name:
            return ""
        if name not in alias:
            alias[name] = f"Customer {len(alias) + 1}"
        return alias[name]
    for i, s in enumerate(fixture["spots"], start=1):
        s["name"] = alias_for(s["name"])
        s["phone"] = f"204-555-{i:04d}" if s["phone"] else ""
        s["plate"] = f"PLT {i:03d}" if s["plate"] else ""
        s["notes"] = ""
        s["updatedBy"] = "fixture"
    for i, rv in enumerate(fixture["reserve"], start=1):
        rv["name"] = alias_for(rv["name"]) or f"Reserve {i}"
        rv["phone"] = f"204-555-9{i:03d}" if rv["phone"] else ""
        rv["plate"] = ""
        rv["notes"] = ""
        rv["updatedBy"] = "fixture"
    fixture["updatedBy"] = "fixture"
    return doc, fixture, oracle


# ------------------------------------------------------------------ mock workbook fixture
def excel_serial(d):
    """Excel 1900-system serial for a date (2026-01-16 -> 46038)."""
    base = datetime(1899, 12, 30)
    if isinstance(d, datetime):
        d = d.date()
    return (datetime(d.year, d.month, d.day) - base).days


def build_workbook_fixture():
    """Anonymised copy of the ledger workbook's data cells, shaped like Graph range values, for tests/unit mocks."""
    wb = load_workbook(LEDGER_XLSX, data_only=True)
    log = wb["LOG"]; srt = wb["LOG_Sorted"]; lists = wb["LISTS"]
    header = [log.cell(row=1, column=c).value for c in range(1, 10)]
    def row_values(ws, r, anon):
        vals = []
        for c in range(1, 10):
            v = ws.cell(row=r, column=c).value
            if c == 3 and v not in (None, ""):
                v = excel_serial(v)
            if c == 7 and v not in (None, ""):
                v = anon.setdefault(str(v).strip(), f"Transaction {len(anon) + 1}")
            if c == 9:
                v = ""
            vals.append("" if v is None else v)
        return vals
    anon = {}
    log_rows = [row_values(log, r, anon) for r in range(2, 28)]          # table body A2:I27 (first row blank)
    sorted_rows = [row_values(srt, r, anon) for r in range(2, 27)]       # LOG_Sorted body A2:I26
    cats = [[lists.cell(row=r, column=c).value if lists.cell(row=r, column=c).value is not None else "" for c in range(1, 5)] for r in range(2, 11)]
    fixture = {
        "source": "TPAHA_2026 (1).xlsx (anonymised: descriptions replaced, notes cleared)",
        "year": int(wb["ConfigHidden"]["A1"].value),
        "priorYearBalance": float(wb["ENTRY"]["B23"].value),
        "logHeader": header,
        "logRows": log_rows,
        "logSortedRows": sorted_rows,
        "categories": cats,
        "numberFormats": {"Timestamp": "yyyy-mm-dd hh:mm:ss", "Date": "yyyy-mm-dd", "Amount": "$#,##0.00"},
        "expectedNextId": max(int(r[0]) for r in log_rows if r[0] not in ("", None)) + 1,
    }
    return fixture


def dump(path, obj):
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(obj, fh, indent=1, ensure_ascii=False)
    print(f"wrote {os.path.relpath(path, ROOT)} ({os.path.getsize(path)} bytes)")


if __name__ == "__main__":
    ledger, ledger_fix, ledger_oracle = build_ledger()
    storage, storage_fix, storage_oracle = build_storage()
    dump(os.path.join(DATA, f"ledger-{ledger['year']}.json"), ledger)
    dump(os.path.join(DATA, "winter-storage.json"), storage)
    dump(os.path.join(FIX, "ledger-fixture.json"), ledger_fix)
    dump(os.path.join(FIX, "ledger-oracle.json"), ledger_oracle)
    dump(os.path.join(FIX, "storage-fixture.json"), storage_fix)
    dump(os.path.join(FIX, "storage-oracle.json"), storage_oracle)
    dump(os.path.join(FIX, "ledger-workbook.json"), build_workbook_fixture())
    print(f"ledger: {len(ledger['transactions'])} transactions, prior balance {ledger['priorYearBalance']}")
    print(f"storage: {sum(1 for s in storage['spots'] if s['name'])} of {len(storage['spots'])} spots assigned, {len(storage['reserve'])} reserve rows")
    notes = [(s['spot'], s['notes']) for s in storage['spots'] if s['notes']]
    print("spot notes present:", notes)
    print("sample oracle spot R1C1:", storage_oracle["spots"]["R1C1"])
    print("footers:", storage_oracle["footers"], "header:", storage_oracle["header"])
